"""Offline checks for PPTX/PPTM conversion and the render script's page handling.

The render test replaces LibreOffice with a stub that copies a prepared PDF, so it
needs Pillow and Poppler (pdftoppm, pdfinfo) but never starts an office suite.
"""
import importlib.util
import io
import os
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover
    Presentation = None
try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
spec = importlib.util.spec_from_file_location("convert_presentation_format", SCRIPTS / "convert_presentation_format.py")
conv = importlib.util.module_from_spec(spec)
spec.loader.exec_module(conv)


def run(argv):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = conv.main([str(a) for a in argv])
    return code, out.getvalue(), err.getvalue()


def add_vba(pptx: Path, pptm: Path) -> None:
    """Turn a PPTX into a PPTM with a fake VBA project, the way Office stores one."""
    with zipfile.ZipFile(pptx) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    ct = parts["[Content_Types].xml"].decode()
    printer = "application/vnd.openxmlformats-officedocument.presentationml.printerSettings"
    other_bins = "".join(f'<Override PartName="/{n}" ContentType="{printer}"/>' for n in parts if n.endswith(".bin"))
    ct = ct.replace(conv.MAIN["pptx"], conv.MAIN["pptm"]).replace(
        f'<Default Extension="bin" ContentType="{printer}"/>', "").replace(
        "</Types>", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
        '<Override PartName="/ppt/vbaData.xml" ContentType="application/vnd.ms-word.vbaData+xml"/>'
        + other_bins + "</Types>")
    assert ct.count('Extension="bin"') == 1
    parts["[Content_Types].xml"] = ct.encode()
    rels = parts["ppt/_rels/presentation.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/'
        'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
    parts["ppt/_rels/presentation.xml.rels"] = rels.encode()
    parts["ppt/vbaProject.bin"] = b"fake vba"
    parts["ppt/_rels/vbaProject.bin.rels"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats'
        b'.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2006/'
        b'relationships/wordVbaData" Target="vbaData.xml"/></Relationships>')
    parts["ppt/vbaData.xml"] = b'<?xml version="1.0"?><wne:vbaSuppData xmlns:wne="urn:x"/>'
    with zipfile.ZipFile(pptm, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)


@unittest.skipIf(Presentation is None, "python-pptx not installed")
class ConvertTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.dir = Path(self.temp.name)
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Hello"
        self.pptx = self.dir / "deck.pptx"
        prs.save(self.pptx)
        self.out = self.dir / "out"

    def parts(self, path):
        with zipfile.ZipFile(path) as z:
            return {n: z.read(n) for n in z.namelist()}

    def test_pptx_to_pptm_changes_only_content_types(self):
        code, out, err = run([self.pptx, "--out-dir", self.out])
        self.assertEqual(code, 0, err)
        before, after = self.parts(self.pptx), self.parts(self.out / "deck.pptm")
        self.assertEqual(set(before), set(after))
        changed = [n for n in before if before[n] != after[n]]
        self.assertEqual(changed, ["[Content_Types].xml"])
        self.assertIn(conv.MAIN["pptm"].encode(), after["[Content_Types].xml"])
        self.assertEqual(Presentation(self.out / "deck.pptm").slides[0].shapes.title.text, "Hello")

    def test_pptm_to_pptx_removes_vba_project(self):
        pptm = self.dir / "macro.pptm"
        add_vba(self.pptx, pptm)
        code, out, err = run([pptm, "--out-dir", self.out])
        self.assertEqual(code, 0, err)
        self.assertIn("vbaProject.bin", out)
        after = self.parts(self.out / "macro.pptx")
        self.assertFalse([n for n in after if "vba" in n.lower()])
        ct = after["[Content_Types].xml"].decode()
        self.assertIn(conv.MAIN["pptx"], ct)
        self.assertNotIn("vbaProject", ct)
        self.assertNotIn(b"vbaProject", after["ppt/_rels/presentation.xml.rels"])
        self.assertEqual(after["ppt/slides/slide1.xml"], self.parts(self.pptx)["ppt/slides/slide1.xml"])
        self.assertEqual(Presentation(self.out / "macro.pptx").slides[0].shapes.title.text, "Hello")

    def test_refuses_existing_target_and_mismatched_extension(self):
        self.assertEqual(run([self.pptx, "--out-dir", self.out])[0], 0)
        code, _, err = run([self.pptx, "--out-dir", self.out])
        self.assertEqual(code, 1)
        self.assertIn("not overwriting", err)
        renamed = self.dir / "fake.pptm"
        shutil.copy(self.pptx, renamed)
        code, _, err = run([renamed, "--out-dir", self.out])
        self.assertEqual(code, 1)
        self.assertIn("not a pptm presentation", err)

    def test_skips_lock_files(self):
        lock = self.dir / "~$deck.pptx"
        lock.write_bytes(b"lock")
        code, _, err = run([lock, "--out-dir", self.out])
        self.assertEqual(code, 0)
        self.assertIn("lock file", err)


@unittest.skipIf(Image is None or not shutil.which("pdftoppm") or not shutil.which("pdfinfo"),
                 "Pillow and Poppler required")
class RenderScriptTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.dir = Path(self.temp.name)
        pages = [Image.new("RGB", (160, 90), (i * 20, 0, 0)) for i in range(12)]
        self.pdf = self.dir / "fixture.pdf"
        pages[0].save(self.pdf, save_all=True, append_images=pages[1:])
        self.stub = self.dir / "soffice"
        self.stub.write_text(
            "#!/bin/sh\n"
            'for a; do case "$prev" in --outdir) out=$a;; esac; prev=$a; last=$a; done\n'
            'case "$1" in -env:UserInstallation=file://*) ;; *) echo "missing private profile" >&2; exit 3;; esac\n'
            f'cp "{self.pdf}" "$out/$(basename "${{last%.*}}").pdf"\n')
        self.stub.chmod(0o755)
        self.deck = self.dir / "deck.pptx"
        self.deck.write_bytes(b"not inspected by the stub")

    def render(self, spec, out):
        env = dict(os.environ, SOFFICE_BIN=str(self.stub), RENDER_DPI="20")
        return subprocess.run(["bash", str(SCRIPTS / "render_slides_to_png.sh"), str(self.deck), spec, str(out)],
                              env=env, capture_output=True, text=True)

    def test_first_pages_on_deck_with_ten_or_more_pages(self):
        r = self.render("2", self.dir / "a")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(sorted(p.name for p in (self.dir / "a").iterdir()), ["deck-第1页.png", "deck-第2页.png"])

    def test_page_range(self):
        r = self.render("9-11", self.dir / "b")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(sorted(p.name for p in (self.dir / "b").iterdir()),
                         ["deck-第10页.png", "deck-第11页.png", "deck-第9页.png"])
        self.assertAlmostEqual(Image.open(self.dir / "b" / "deck-第10页.png").getpixel((5, 5))[0], 9 * 20, delta=4)

    def test_rejects_out_of_range_and_existing_outputs(self):
        self.assertNotEqual(self.render("12-13", self.dir / "c").returncode, 0)
        self.assertEqual(self.render("1", self.dir / "d").returncode, 0)
        r = self.render("1", self.dir / "d")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("已存在", r.stderr)


if __name__ == "__main__":
    unittest.main()
