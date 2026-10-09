#!/usr/bin/env python3
"""Read-only XLSX attendance audit; emit JSON to stdout, never rewrite evidence."""
from __future__ import annotations

import argparse
import itertools
import json
import re
import sys
import unicodedata
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path


def normalize(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value)))


HEADERS = {
    "user_id": normalize("用户ID"),
    "name": normalize("姓名"),
    "duration": normalize("在线时长（分钟）"),
}
OPTIONAL_HEADERS = {
    "hospital": normalize("医院"),
    "department": normalize("科室"),
}


def parse_minutes(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    text = normalize(value)
    if text.endswith("分钟"):
        text = text[:-2]
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() and number >= 0 else None


def find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:30]):
        labels = [normalize(cell) for cell in row]
        if all(label in labels for label in HEADERS.values()):
            if any(labels.count(label) != 1 for label in HEADERS.values()):
                raise ValueError("参会表必需列存在重复表头，需明确列映射")
            columns = {key: labels.index(label) for key, label in HEADERS.items()}
            for key, label in OPTIONAL_HEADERS.items():
                if label in labels:
                    if labels.count(label) != 1:
                        raise ValueError(f"展示列存在重复表头：{label}")
                    columns[key] = labels.index(label)
            return index, columns
    return None


def parse_rows(rows: list[tuple], source: str, sheet: str, meeting: str) -> dict:
    header = find_header(rows)
    if header is None:
        raise ValueError("前30行未找到用户ID/姓名/在线时长（分钟）表头")
    header_index, columns = header
    records, issues = [], []
    for index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        def cell(key: str) -> object:
            column = columns.get(key)
            return row[column] if column is not None and column < len(row) else None

        # Ignore completely blank table rows, not partially populated participants.
        if not any(normalize(cell(key)) for key in HEADERS):
            continue
        if all(normalize(cell(key)) == label for key, label in HEADERS.items()):
            continue  # repeated print header, not a participant
        key = tuple(normalize(cell(field)) for field in ("user_id", "name"))
        minutes = parse_minutes(cell("duration"))
        records.append({
            "row": index, "key": key if all(key) else None,
            "user_id": key[0], "name": key[1],
            "hospital": normalize(cell("hospital")), "department": normalize(cell("department")),
            "reported_minutes": str(minutes) if minutes is not None else None,
        })
        if not all(key):
            issues.append({"type": "incomplete_identity", "rows": [index]})
        if minutes is None:
            issues.append({"type": "unknown_duration", "rows": [index]})
    return {
        "source": source, "sheet": sheet, "meeting": meeting,
        "header_row": header_index + 1, "records": records, "data_issues": issues,
    }


def summarize(report: dict, min_minutes: Decimal, min_people: int) -> dict:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    issues = list(report["data_issues"])
    unidentified = []
    for record in report["records"]:
        if record["key"] is None:
            unidentified.append(record)
        else:
            groups[tuple(record["key"])].append(record)
    id_names: dict[str, set[str]] = defaultdict(set)
    for user_id, name in groups:
        id_names[user_id].add(name)
    conflicting_ids = {user_id for user_id, names in id_names.items() if len(names) > 1}
    for user_id in sorted(conflicting_ids):
        issues.append({
            "type": "id_name_conflict", "user_id": user_id,
            "rows": [record["row"] for record in report["records"] if record["user_id"] == user_id],
        })
    eligible, unknown, below = [], [], []
    members = []
    for key, records in sorted(groups.items()):
        values = {record["reported_minutes"] for record in records}
        rows = [record["row"] for record in records]
        members.append({"key": key, "rows": rows})
        if len(records) > 1:
            issues.append({"type": "duplicate_identity", "rows": rows})
        if key[0] in conflicting_ids:
            unknown.append(key)
        elif len(values) != 1:
            issues.append({"type": "conflicting_duplicate_duration", "rows": rows})
            unknown.append(key)
        elif None in values:
            unknown.append(key)
        elif Decimal(next(iter(values))) >= min_minutes:
            eligible.append(key)
        else:
            below.append(key)
    # Partial identities cannot establish distinct people; each row adds only to an upper bound.
    possible_unidentified = sum(
        record["reported_minutes"] is None or Decimal(record["reported_minutes"]) >= min_minutes
        for record in unidentified
    )
    lower = len(eligible)
    upper = lower + len(unknown) + possible_unidentified
    status = "pass" if lower >= min_people else "fail" if upper < min_people else "needs_review"
    return {
        **report, "data_issues": issues, "record_count": len(report["records"]),
        "unique_participant_count": len(groups), "identity_complete": not unidentified and not conflicting_ids,
        "members": members, "eligible_count": lower, "possible_eligible_count": upper,
        "unknown_eligibility_count": len(unknown) + possible_unidentified,
        "below_threshold_count": len(below), "headcount_status": status,
    }


def compare_pairs(reports: list[dict], policy: str, threshold: Decimal) -> list[dict]:
    pairs = []
    for a, b in itertools.combinations(reports, 2):
        map_a = {tuple(member["key"]): member["rows"] for member in a["members"]}
        map_b = {tuple(member["key"]): member["rows"] for member in b["members"]}
        common = sorted(map_a.keys() & map_b.keys())
        size_a, size_b, count = len(map_a), len(map_b), len(common)
        ratio_a = Decimal(count) / size_a if size_a else None
        ratio_b = Decimal(count) / size_b if size_b else None
        names_a: dict[str, set[str]] = defaultdict(set)
        names_b: dict[str, set[str]] = defaultdict(set)
        for user_id, name in map_a:
            names_a[user_id].add(name)
        for user_id, name in map_b:
            names_b[user_id].add(name)
        conflicts = sorted(user_id for user_id in names_a.keys() & names_b.keys() if names_a[user_id] != names_b[user_id])
        complete = a["identity_complete"] and b["identity_complete"] and size_a > 0 and size_b > 0 and not conflicts
        if complete:
            denominator = min(size_a, size_b) if policy == "larger-ratio" else max(size_a, size_b)
            # Compare unrounded exact inputs; display rounding cannot change the verdict.
            failed = Decimal(count) > threshold * denominator
            status = "fail" if failed else "pass"
        else:
            status = "needs_review"
        pairs.append({
            "meeting_a": a["meeting"], "meeting_b": b["meeting"],
            "source_a": a["source"], "source_b": b["source"],
            "size_a": size_a, "size_b": size_b, "common_count": count,
            "ratio_a": str(ratio_a) if ratio_a is not None else None,
            "ratio_b": str(ratio_b) if ratio_b is not None else None,
            "percent_a": f"{ratio_a * 100:.2f}%" if ratio_a is not None else None,
            "percent_b": f"{ratio_b * 100:.2f}%" if ratio_b is not None else None,
            "policy": policy, "status": status,
            "id_name_conflicts": conflicts,
            "matched_members": [{"key": key, "rows_a": map_a[key], "rows_b": map_b[key]} for key in common],
        })
    return pairs


def meeting_label(path: Path) -> str:
    parent = path.parent
    if parent.name in {"参会者签到表", "参会明细", "签到表"}:
        parent = parent.parent
    return parent.name or path.stem


def read_report(path: Path, sheet_name: str | None = None) -> dict:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name:
            if sheet_name not in book.sheetnames:
                raise ValueError(f"未找到指定工作表：{sheet_name}")
            sheets = [book[sheet_name]]
        elif "参会报告" in book.sheetnames:
            sheets = [book["参会报告"]]
        else:
            sheets = book.worksheets
        candidates = []
        for sheet in sheets:
            iterator = sheet.iter_rows(values_only=True)
            prefix = list(itertools.islice(iterator, 30))
            if find_header(prefix) is not None:
                rows = prefix + list(iterator)
                candidates.append(parse_rows(rows, str(path), sheet.title, meeting_label(path)))
        if len(candidates) != 1:
            raise ValueError(f"找到 {len(candidates)} 个可识别参会工作表；需核对表头或用 --sheet 选择")
        return candidates[0]
    finally:
        book.close()


def discover(inputs: list[str], pattern: str) -> tuple[list[Path], list[dict]]:
    files, errors = set(), []
    for text in inputs:
        path = Path(text).expanduser().resolve()
        if path.is_dir():
            found = [item for item in path.rglob(pattern) if item.is_file() and not item.name.startswith("~$")]
            if not found:
                errors.append({"source": str(path), "error": f"未发现匹配 {pattern} 的文件"})
            files.update(found)
        elif path.is_file() and not path.name.startswith("~$"):
            files.add(path)
        else:
            errors.append({"source": str(path), "error": "路径不存在或是 Excel 锁文件"})
    return sorted(files), errors


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", help="会议批次目录或多个参会明细xlsx文件")
    parser.add_argument("--pattern", default="*参会明细*.xlsx")
    parser.add_argument("--sheet")
    parser.add_argument("--min-minutes", type=Decimal, default=Decimal(20))
    parser.add_argument("--min-people", type=int, default=10)
    parser.add_argument("--max-overlap", type=Decimal, default=Decimal("0.5"))
    parser.add_argument("--overlap-policy", choices=["larger-ratio", "smaller-ratio"], default="larger-ratio")
    args = parser.parse_args(argv)
    if not args.min_minutes.is_finite() or args.min_minutes < 0 or args.min_people < 1:
        parser.error("时长阈值必须是非负有限数字，人数阈值必须大于0")
    if not args.max_overlap.is_finite() or not 0 <= args.max_overlap <= 1:
        parser.error("重合率阈值必须是0到1之间的有限数字")
    files, errors = discover(args.inputs, args.pattern)
    reports = []
    for path in files:
        try:
            if path.suffix.lower() != ".xlsx":
                raise ValueError("仅支持xlsx；不转换或改写原件")
            reports.append(summarize(read_report(path, args.sheet), args.min_minutes, args.min_people))
        except Exception as exc:
            errors.append({"source": str(path), "error": str(exc)})
    pairs = compare_pairs(reports, args.overlap_policy, args.max_overlap)
    labels = [report["meeting"] for report in reports]
    duplicate_labels = sorted(label for label in set(labels) if labels.count(label) > 1)
    warnings = (["存在相同会议目录名，请按source核实是否重复版本：" + ", ".join(duplicate_labels)] if duplicate_labels else [])
    result = {
        "rules": {
            "duration_source": "在线时长（分钟）", "min_minutes_inclusive": str(args.min_minutes),
            "min_people_inclusive": args.min_people, "max_overlap_inclusive": str(args.max_overlap),
            "identity_key": ["用户ID", "姓名"], "overlap_population": "完整去重名单",
            "overlap_policy": args.overlap_policy,
            "overlap_policy_note": "分母是可调整的默认约定，不是已核实的平台政策",
        },
        "discovered_files": [str(path) for path in files], "reports": reports,
        "pair_count": len(pairs), "expected_pair_count": len(reports) * (len(reports) - 1) // 2,
        "pairs": pairs, "errors": errors, "warnings": warnings,
        "headcount_failure_count": sum(report["headcount_status"] == "fail" for report in reports),
        "overlap_failure_count": sum(pair["status"] == "fail" for pair in pairs),
        "screenshots_checked": False,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    # Findings are data, not execution failures; unreadable/missing inputs return a nonzero code.
    return 2 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
